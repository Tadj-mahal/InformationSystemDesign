# Import the necessary modules
from pygismeteo import Gismeteo   # For accessing weather data from Gismeteo
import pandas as pd               # For data manipulation and analysis
import tkinter as tk              # For building GUI applications
import openpyxl                   # For working with Excel files
import xlsxwriter                 # For creating Excel charts
import datetime                   # For working with dates and times

# Define a function to retrieve the weather information for a given city


def get_weather(city):
    # Create a Gismeteo object
    gismeteo = Gismeteo()
    # Search for the city by name and get its ID
    search_results = gismeteo.search.by_query(city)
    city_id = search_results[0].id
    # Get the current weather information for the city
    current = gismeteo.current.by_id(city_id)
    # Extract the temperature and weather condition from the data
    temperature = current.temperature.air.c
    condition = current.description.full
    # Return the temperature and condition
    return temperature, condition

# Define a function to update the weather information displayed in the GUI


def update_weather():
    # Get the city name from the entry field
    city = city_entry.get()
    # Get the current temperature and weather condition for the city
    temperature, condition = get_weather(city)
    # Get the current date and time
    current_datetime = datetime.datetime.now()
    # Format the date and time string
    formatted_datetime = current_datetime.strftime("%Y-%m-%d %H:%M:%S")
    # Update the temperature label with the current temperature
    temperature_label.config(text=f"Temperature: {temperature}°C")
    # Update the condition label with the current weather condition
    condition_label.config(text=f"Condition: {condition}")
    # Update the current date and time label with the formatted date and time string
    current_data.config(text=f"Current date and time: {formatted_datetime}")

# Define a function to add the current weather information to an Excel file


def add_in_excel():
    # Get the city name from the entry field
    city = city_entry.get()
    # Get the current temperature and weather condition for the city
    temperature, condition = get_weather(city)
    # Get the current date and time
    current_datetime = datetime.datetime.now()
    # Format the date and time string
    formatted_datetime = current_datetime.strftime("%Y-%m-%d %H:%M:%S")

    # Specify the filename and sheetname for the Excel file
    filename = 'weather.xlsx'
    sheetname = 'weather1'
    # Read the existing data from the Excel file into a Pandas DataFrame
    df = pd.read_excel(filename, sheet_name=sheetname)
    # Create a new row for the current weather information
    new_row = pd.DataFrame([[city, temperature, condition, current_datetime]],
                           index=[city], columns=['city', 'temperature', 'condition', 'data'])
    # Append the new row to the existing data
    df = df.append(new_row, ignore_index=False)
    # Write the updated data to the Excel file
    df.to_excel(filename, sheet_name=sheetname, index=False)
    # Add a chart to the Excel file
    add_chart(filename, sheetname)
    # Update the label to indicate that the Excel file was saved successfully
    save_excel.config(text=f"Successfull saved Excel-file")

# Define a function to add a chart to an Excel file


def add_chart(filename, sheetname):
    # Read the data from the Excel file into a Pandas DataFrame
    df = pd.read_excel(filename)

    # Create a bar plot of the temperature data for each city
    ax = df.plot.bar(x='city', y='temperature')
    # Add labels for each bar indicating the temperature value
    for i, v in enumerate(df['temperature']):
        ax.text(i, v + 1, str(v), ha='center')

    # Write the DataFrame back to the Excel file
    writer = pd.ExcelWriter(filename, engine='xlsxwriter')
    df.to_excel(writer, sheet_name=sheetname, index=False)
    workbook = writer.book
    worksheet = writer.sheets[sheetname]

    # Create a new column chart object
    chart = workbook.add_chart({'type': 'column'})
    # Add a data series to the chart for the temperature data
    chart.add_series({
        'name': 'Temperature',
        'categories': '=weather1!$A$2:$A$10',
        'values': '=weather1!$B$2:$B$10',
    })

    # Insert the chart into the worksheet at cell G2
    worksheet.insert_chart('G2', chart)
    # Save the changes to the Excel file
    writer.save()


# create window and set its title
window = tk.Tk()
window.title("Weather App")

# create label for city input field
city_label = tk.Label(window, text="Enter city name:")
city_label.pack()

# create input field for city name
city_entry = tk.Entry(window)
city_entry.pack()

# create button to update weather information
update_button = tk.Button(
    window, text="Update Weather", command=update_weather)
update_button.pack()

# create button to add weather information to Excel file and display a chart
excel_button = tk.Button(window, text="Add in Excel", command=add_in_excel)
excel_button.pack()

# create label to display temperature information
temperature_label = tk.Label(window, text="Temperature: ")
temperature_label.pack()

# create label to display weather condition information
condition_label = tk.Label(window, text="Condition: ")
condition_label.pack()

# create label to display current date and time information
current_data = tk.Label(window, text="Current date and time: ")
current_data.pack()

# create label to display whether the Excel file has been saved or not
save_excel = tk.Label(window, text="Excel-file is not saved")
save_excel.pack()

# start the main event loop for the window
window.mainloop()
